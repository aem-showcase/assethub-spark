#!/usr/bin/env python3
"""Generate DA-ready .docx files for the Disney India ('disney-in') demo portal.

Mirrors docs/da-content/create-docs.py's authoring conventions (same block
tables, section metadata, hyperlink helpers) but scoped ONLY to the company
folder /companies/disney-in, per Step 4.3 of the rebrand-portal skill:
  - Rewrites real Disney India copy (not a name-swap of the base template).
  - Swaps the logo shortcode :frescopa-icon:/:frescopa-beans: ->
    :disney-in-icon:/:disney-in-beans:.
  - Scopes every internal link to /companies/disney-in/... EXCEPT the
    welcome page's Sign in link, which MUST stay /auth/login (the worker's
    auth routes are not prefixed with DEMO_BASE_PATH).
  - Category contract derived from disney.in's real navigation/sections:
    Movies, Marvel, Disney Parks, Disney Cruise, Classic Animation.

IMPORTANT (DA workflow, same as the base script):
  Upload each .docx DIRECTLY into DA (drag into the DA file browser) at the
  path printed after each Wrote line. Do NOT open/save these files in
  Microsoft Word first -- Word rewrites the relative links as file:// paths
  and collapses the section breaks, which breaks the pages.

This script performs NO network calls -- it only writes local .docx files
under migration-work/da-content/output/. Uploading them to DA (drag-drop or
an authorized publish path) is a separate, explicit step.
"""

from pathlib import Path
from urllib.parse import quote
import json

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = Path(__file__).parent / 'output'
OUT.mkdir(parents=True, exist_ok=True)

COMPANY_KEY = 'disney-in'
COMPANY_ROOT = f'/companies/{COMPANY_KEY}'
ICON = f':{COMPANY_KEY}-icon:'
BEANS = f':{COMPANY_KEY}-beans:'

ASSET_EXC_FACETS = """{
"productCategory": {
"label": "Category",
"type": "string"
},
"campaign": {
"label": "Campaign",
"type": "string"
},
"channel": {
"label": "Channel",
"type": "string"
},
"dc:subject": {
"label": "Keywords",
"type": "string"
},
"dc:format": {
"label": "Format",
"type": "string"
}
}"""

BRAND_STORY = [
    'Disney India brings the magic of Disney, Marvel, Pixar, and Star Wars '
    'storytelling to audiences across the country \u2014 in theatres, on Disney+ '
    'Hotstar, and through Disney Parks & Cruise experiences worldwide.',
    'This portal is the home of the Disney India brand asset library \u2014 search, '
    'preview, and download approved film, character, park, and cruise imagery '
    'for marketing and partner use.',
]

CATEGORY_FACET_KEY = 'productCategory'

# Derived from disney.in's real navigation/sections (Step 4 preflight).
CATEGORY_TILES = [
    ('Movies', 'Latest theatrical releases and Disney+ Hotstar originals.', 'movies'),
    ('Marvel', 'Marvel Cinematic Universe characters, posters, and key art.', 'marvel'),
    ('Disney Parks', 'Disneyland and Walt Disney World imagery and experiences.', 'disney-parks'),
    ('Disney Cruise', 'Disney Cruise Line ships, itineraries, and onboard moments.', 'disney-cruise'),
    ('Classic Animation', 'Timeless Disney animated classics and character art.', 'classic-animation'),
]


def company_path(rel):
    """Prefix an internal relative path with the company folder."""
    rel = rel.lstrip('/')
    return f'{COMPANY_ROOT}/{rel}'


def category_search_url(facet_value):
    facet_filters = json.dumps(
        {CATEGORY_FACET_KEY: {facet_value: True}},
        separators=(',', ':'),
    )
    return company_path(f'en/search?facetFilters={quote(facet_filters)}')


def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'auto')
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_hyperlink(paragraph, text, url, bold=False, italic=False):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    if bold:
        rPr.append(OxmlElement('w:b'))
    if italic:
        rPr.append(OxmlElement('w:i'))
    run.append(rPr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_button(doc, text, url, secondary=False):
    p = doc.add_paragraph()
    add_hyperlink(p, text, url, bold=not secondary, italic=secondary)
    return p


def add_metadata_table(doc, rows):
    table = doc.add_table(rows=1 + len(rows), cols=2)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = 'Metadata'
    for i, (a, b) in enumerate(rows, start=1):
        table.rows[i].cells[0].text = a
        table.rows[i].cells[1].text = b


def add_section_metadata(doc, style):
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = 'Section Metadata'
    table.rows[1].cells[0].text = 'style'
    table.rows[1].cells[1].text = style
    doc.add_paragraph('')


def add_nav_role_metadata(doc, role):
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = 'Section Metadata'
    table.rows[1].cells[0].text = 'role'
    table.rows[1].cells[1].text = role
    doc.add_paragraph('')


def add_block_table(doc, block_name, rows=None, cols=2):
    rows = rows if rows is not None else [['', '']]
    table = doc.add_table(rows=1 + len(rows), cols=cols)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = block_name
    for i, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.rows[i].cells[c].text = value
    doc.add_paragraph('')


def add_category_carousel(doc, cards, block_name='carousel (tiles)'):
    table = doc.add_table(rows=1 + len(cards), cols=2)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = block_name
    for i, (title, blurb, facet_value) in enumerate(cards, start=1):
        table.rows[i].cells[0].text = ''  # image cell, filled per Step 5 asset enrichment
        cell = table.rows[i].cells[1]
        cell.paragraphs[0].text = ''
        head = cell.paragraphs[0]
        head.style = doc.styles['Heading 3']
        head.add_run(title)
        cell.add_paragraph(blurb)
        link_p = cell.add_paragraph()
        add_hyperlink(link_p, 'Browse \u2192', category_search_url(facet_value), bold=True)
    doc.add_paragraph('')


def build_index():
    """Landing: branded hero (search bar) + Browse-by-category."""
    doc = Document()

    doc.add_paragraph(ICON)
    doc.add_heading('Find your Disney India assets', level=1)
    doc.add_paragraph(
        'Search, preview, and download approved Disney, Marvel, Pixar, and '
        'Star Wars imagery for India marketing and partner use.'
    )
    add_block_table(doc, 'search-bar', [
        ('enableSemanticSearch', 'true'),
    ])
    add_section_metadata(doc, 'search-hero')
    add_hr(doc)

    doc.add_heading('Browse by category', level=2)
    add_category_carousel(doc, CATEGORY_TILES)
    add_section_metadata(doc, 'category-tiles')
    add_hr(doc)

    add_metadata_table(doc, [
        ('title', 'Disney India Asset Library'),
        ('description', 'Search, preview, and download approved Disney India brand assets.'),
    ])

    path = OUT / 'index.docx'
    doc.save(path)
    print(f'Wrote {path}  -> upload into DA: aem-showcase/assethub-spark -> {company_path("en/index")}')


def build_about():
    doc = Document()

    doc.add_heading('About Disney India', level=1)
    for para in BRAND_STORY:
        doc.add_paragraph(para)
    add_button(doc, 'Browse the library', company_path('en/search'))
    add_hr(doc)
    add_metadata_table(doc, [
        ('title', 'About Disney India'),
        ('description', 'The story behind the Disney India brand asset library.'),
    ])

    path = OUT / 'about.docx'
    doc.save(path)
    print(f'Wrote {path}  -> upload into DA: aem-showcase/assethub-spark -> {company_path("en/about")}')


def build_nav():
    doc = Document()

    brand = doc.add_paragraph()
    add_hyperlink(brand, ICON, company_path('en/'))
    add_nav_role_metadata(doc, 'brand')
    add_hr(doc)

    doc.add_paragraph('\u00a0')
    add_nav_role_metadata(doc, 'tools')

    path = OUT / 'nav.docx'
    doc.save(path)
    print(f'Wrote {path}  -> upload into DA: aem-showcase/assethub-spark -> {company_path("en/nav")}')


def build_footer():
    doc = Document()

    doc.add_paragraph(ICON)
    doc.add_paragraph('The magic of Disney, Marvel, Pixar, and Star Wars \u2014 in India.')
    doc.add_paragraph('Your brand asset distribution portal.')
    add_hr(doc)

    columns = [
        ('Assets', [
            ('Search Assets', company_path('en/search')),
            ('Browse Collections', company_path('en/search-collections')),
        ]),
        ('Company', [
            ('About', company_path('en/about')),
            ('Contact', 'mailto:assets@disney-in-demo.example'),
        ]),
        ('Help', [
            ('Documentation', company_path('en/about')),
            ('Support', 'mailto:assets@disney-in-demo.example'),
        ]),
    ]
    table = doc.add_table(rows=2, cols=len(columns))
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = 'columns'
    for i, (header, links) in enumerate(columns):
        cell = table.rows[1].cells[i]
        cell.paragraphs[0].text = ''
        head = cell.paragraphs[0]
        head.add_run(header).bold = True
        for label, url in links:
            link_p = cell.add_paragraph()
            add_hyperlink(link_p, label, url)
    doc.add_paragraph('')
    add_hr(doc)

    doc.add_paragraph('\u00a9 2026 Disney India. All rights reserved.')

    path = OUT / 'footer.docx'
    doc.save(path)
    print(f'Wrote {path}  -> upload into DA: aem-showcase/assethub-spark -> {company_path("en/footer")}')


def build_welcome():
    """Login/welcome page: shown to unauthenticated users at
    /companies/disney-in/public/welcome.

    Sign in link MUST stay /auth/login (NOT company-scoped) -- the worker's
    auth routes are hardcoded at /auth/* and are not prefixed with
    DEMO_BASE_PATH (see cloudflare/src/auth.js AUTH_PREFIX). A scoped link
    here causes an unauthenticated redirect loop.
    """
    doc = Document()

    doc.add_paragraph(BEANS)

    doc.add_heading('Welcome to the Disney India Asset Library', level=1)
    doc.add_paragraph(
        'Search, preview, and download approved Disney, Marvel, Pixar, and '
        'Star Wars imagery for India marketing and partner use.'
    )

    btn_p = doc.add_paragraph()
    add_hyperlink(btn_p, 'Sign in', '/auth/login', bold=True)

    add_section_metadata(doc, 'welcome')
    add_hr(doc)

    add_metadata_table(doc, [
        ('title', 'Sign in \u2014 Disney India Asset Library'),
        ('description', 'Sign in to access the Disney India brand asset portal.'),
        ('header', 'no'),
        ('footer', 'no'),
    ])

    path = OUT / 'welcome.docx'
    doc.save(path)
    print(f'Wrote {path}  -> upload into DA: aem-showcase/assethub-spark -> {company_path("public/welcome")}')


def main():
    build_index()
    build_about()
    build_nav()
    build_footer()
    build_welcome()
    print(
        '\nDone. These 5 .docx files are the Step 4.3 content-register rewrite for '
        f'{COMPANY_ROOT}, scoped ONLY to that folder (shared-root /en/nav, /en/footer '
        'are untouched).'
        '\n\nUpload each .docx DIRECTLY into https://da.live/edit#/aem-showcase/assethub-spark'
        f'{COMPANY_ROOT}/... at the path printed after each "Wrote" line above'
        '\n(do NOT open/save them in Microsoft Word first -- this rewrites relative links as'
        '\nfile:// paths and collapses section breaks).'
        '\n\nAfter upload, DA auto-saves each as a preview-ready source doc. Then run'
        '\nadmin.hlx.page preview + live for each path to publish (Step 4.4).'
    )


if __name__ == '__main__':
    main()
